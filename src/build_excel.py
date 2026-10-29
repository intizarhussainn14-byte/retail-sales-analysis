"""Create the seven-sheet Excel analysis from verified outputs, without re-cleaning.

Authoring: openpyxl. Native calculation, PivotTables and PDF previews: installed
Microsoft Excel via finalize_excel.ps1. Requires Windows and Excel Desktop.
"""
import argparse
import json
import os
import shutil
import subprocess
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from inspect_data import ROOT, sha256
from project_paths import project_path
from reconcile_sql import KEYS, compare_tables

GBP = '"£"#,##0.00;[Red]("£"#,##0.00)'
INTEGER, PERCENT = '#,##0', '0.00%;[Red](0.00%)'
NAVY, TEAL, LIGHT, AMBER = '17324D', '087E8B', 'EAF3F6', 'FFF1CC'
SHEETS = ['Overview', 'Monthly Sales', 'Top Products', 'Country Performance',
          'Customer Analysis', 'Data Quality', 'Sources']


def build(args):
    analysis = args.analysis.resolve()
    metadata = json.loads((analysis/'analysis_metadata.json').read_text())
    cleaned = project_path(metadata['input'], 'data/processed')
    cleaning_file = project_path(metadata['cleaning_summary'], 'reports')
    cleaning = json.loads(cleaning_file.read_text())
    if sha256(cleaned) != metadata['input_sha256']:
        raise ValueError('Cleaned data no longer matches the verified analysis')
    evidence = json.loads(args.sql_evidence.read_text())
    assert evidence['status'] == 'PASS'
    assert project_path(evidence['python_report'], 'reports') == analysis
    frames = {n: pd.read_csv(analysis/f'{n}.csv', dtype='string', keep_default_na=False) for n in KEYS}
    sql_folder = ROOT/evidence['results_folder']
    for name, frame in frames.items():
        actual = sql_folder/f'{name}.tsv'
        if actual.exists():
            assert sha256(actual) == evidence['tables'][name]['result_sha256']
            compare_tables(name, frame, pd.read_csv(actual, sep='\t', dtype='string', keep_default_na=False))
    kpi, repeat = frames['kpis'].iloc[0], frames['repeat_purchasing'].iloc[0]
    customers = frames['customer_purchases'].copy()
    customers['CompletedOrders'] = customers.CompletedOrders.astype(int)
    customers['RevenueGBP'] = customers.RevenueGBP.map(Decimal)
    assert len(customers) == int(kpi.KnownUniqueCustomers)
    assert customers.CompletedOrders.ge(2).sum() == int(repeat.RepeatCustomers)
    assert sum(Decimal(x) for x in frames['monthly'].RevenueGBP) == Decimal(kpi.RevenueGBP)
    assert sum(Decimal(x) for x in frames['countries'].RevenueGBP) == Decimal(kpi.RevenueGBP)
    assert cleaning['completed_sales_GBP'] == kpi.RevenueGBP
    assert sum(int(x) for x in frames['monthly'].CompletedOrders) == int(kpi.CompletedOrders)
    run = datetime.now(timezone.utc).strftime('run_%Y%m%dT%H%M%S_%fZ')
    work = ROOT/'reports/excel_qa'/run
    work.mkdir(parents=True, exist_ok=False)
    book = Workbook()
    book.remove(book.active)
    expected = []
    references = {}

    def expect(sheet, address, value, kind='number'):
        expected.append({'sheet': sheet.title, 'cell': address, 'expected': str(value), 'kind': kind})

    def put(sheet, address, value, fmt=None, kind=None):
        sheet[address] = value
        if fmt:
            sheet[address].number_format = fmt
        if kind:
            expect(sheet, address, value, kind)

    def note(sheet, row, text, end='G', height=32, fill=None):
        sheet.merge_cells(f'A{row}:{end}{row}')
        c = sheet.cell(row, 1, text)
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.font = Font(name='Calibri', size=11, color=NAVY)
        if fill:
            c.fill = PatternFill('solid', fgColor=fill)
        sheet.row_dimensions[row].height = height

    def section(sheet, row, text, end='G'):
        note(sheet, row, text, end, 25, NAVY)
        sheet.cell(row, 1).font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')

    def table(sheet, row, headers, rows, name):
        for c, header in enumerate(headers, 1):
            sheet.cell(row, c, header)
        for r, values in enumerate(rows, row+1):
            for c, value in enumerate(values, 1):
                sheet.cell(r, c, value)
                sheet.cell(r, c).alignment = Alignment(vertical='center', wrap_text=c==1)
        ref = f'A{row}:{get_column_letter(len(headers))}{row+len(rows)}'
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sheet.add_table(t)
        for c in sheet[row][:len(headers)]:
            c.font = Font(name='Calibri', color='FFFFFF', bold=True, size=10)
            c.fill = PatternFill('solid', fgColor=NAVY)
            c.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[row].height = 32
        return row+1, row+len(rows)

    for name in SHEETS:
        s = book.create_sheet(name)
        s.sheet_view.showGridLines = False
        s.sheet_view.zoomScale = 85
        s.freeze_panes = 'B6'
        for col, width in {'A':32,'B':21,'C':18,'D':22,'E':20,'F':19,'G':29,'H':3,
                           'I':19,'J':21,'K':18,'L':16}.items():
            s.column_dimensions[col].width = width
        s.sheet_properties.pageSetUpPr.fitToPage = True
        s.page_setup.orientation = 'landscape'
        s.page_setup.paperSize = s.PAPERSIZE_A3
        s.page_setup.fitToWidth, s.page_setup.fitToHeight = 1, 0
        s.print_options.horizontalCentered = True
        s.page_margins.left = s.page_margins.right = .25
        s.page_margins.top = s.page_margins.bottom = .35
        s.print_title_rows = '1:5'
        s.oddFooter.center.text = 'Retail Sales Analysis | &A | Page &P of &N'
        s.oddFooter.center.size = 9
        s.sheet_properties.tabColor = TEAL
        for row in range(1, 82):
            s.row_dimensions[row].height = 20
        s.merge_cells('A1:G2')
        s['A1'] = name
        s['A1'].font = Font(name='Calibri', size=24, bold=True, color='FFFFFF')
        s['A1'].fill = PatternFill('solid', fgColor=NAVY)
        s['A1'].alignment = Alignment(vertical='center')
        s.row_dimensions[1].height = 28
        s.row_dimensions[2].height = 14

    sources = book['Sources']
    note(sources, 3, 'Traceable snapshot of verified Python outputs and MariaDB reconciliation. No cleaning rules were changed.', height=30)
    section(sources, 5, 'Dataset attribution and analytical lineage')
    source_notes = [
        'Chen, D. (2015). Online Retail [Dataset]. UCI Machine Learning Repository.',
        'Dataset: https://archive.ics.uci.edu/dataset/352/online+retail',
        'DOI: https://doi.org/10.24432/C5BW33 | License: https://creativecommons.org/licenses/by/4.0/',
        'CC BY 4.0: give credit, link the license, indicate modifications and do not imply endorsement. Cleaning and aggregation are project modifications.',
        'Observed coverage: 01 Dec 2010 08:26 to 09 Dec 2011 12:50. Dec 2010 begins on day 1 but completeness is unverified; Dec 2011 is partial.',
        'Analysis input: '+analysis.relative_to(ROOT).as_posix(),
        'Cleaning evidence: '+cleaning_file.relative_to(ROOT).as_posix(),
        'MariaDB 10.4.32: all seven equivalent SQL tables matched Python. See reports/sql_reconciliation.md.',
        'Cleaned input SHA-256: '+metadata['input_sha256'],
        'Raw SHA-256: '+cleaning['source_sha256'],
        'Historical analysis files may say SQL was unexecuted at generation; the later SQL reconciliation report is the current evidence.'
    ]
    for row, text in enumerate(source_notes, 6):
        note(sources, row, text, height=30)
    ref_values = [('Revenue GBP', Decimal(kpi.RevenueGBP), GBP), ('Completed orders', int(kpi.CompletedOrders), INTEGER),
                  ('Units sold', int(kpi.UnitsSold), INTEGER), ('Average order value GBP', Decimal(kpi.AverageOrderValueGBP), GBP),
                  ('Known customers', int(kpi.KnownUniqueCustomers), INTEGER), ('Repeat customers', int(repeat.RepeatCustomers), INTEGER),
                  ('Repeat rate (rounded reference)', Decimal(repeat.RepeatCustomerPct)/100, PERCENT),
                  ('Retained missing-customer lines', cleaning['retained_missing_customer_rows'], INTEGER),
                  ('Raw rows', cleaning['input_rows'], INTEGER), ('Completed-sales rows', cleaning['completed_sales_rows'], INTEGER),
                  ('Excluded rows', cleaning['excluded_rows'], INTEGER)]
    table(sources, 19, ['Verified reference metric','Python value','Unit / role'],
          [[n,v,'GBP' if fmt==GBP else ('Rate' if fmt==PERCENT else 'Count')] for n,v,fmt in ref_values], 'VerifiedReferences')
    for row, (name, value, fmt) in enumerate(ref_values,20):
        sources.cell(row,2).number_format = fmt
        references[name] = f"'Sources'!$B${row}"
        expect(sources,f'B{row}',value,'percent' if fmt==PERCENT else 'number')
    section(sources, 33, 'Regenerate and refresh')
    for row, text in enumerate([
        r'From the project root: .\.venv\Scripts\python.exe src/build_excel.py',
        r'To regenerate an existing workbook safely: .\.venv\Scripts\python.exe src/build_excel.py --replace',
        'The --replace option archives the existing workbook before replacement. Requirements: Python environment plus installed Windows Microsoft Excel for calculation, native pivots and preview export.',
        'For a new verified analysis: pass --analysis reports/ANALYSIS_RUN and --sql-evidence reports/SQL_RECONCILIATION.json. They must refer to the same data.',
        'Refresh All updates embedded PivotTable caches only. It does not rerun Python or fetch CSV changes. Regenerate to replace the snapshot and rerun validation.',
        'Summary table filters and PivotTable filters do not recalculate the full-dataset Overview cards. Charts reference their nearby summary data; restore filters before reconciliation.',
        'See reports/excel_refresh_and_pivottables.md for exact transaction-level Power Query/Data Model steps. Distinct invoices are essential: ordinary Count of invoice lines is not order count.'
    ],34):
        note(sources,row,text,height=38)
    section(sources,43,'Native PivotTables and manual recreation')
    for row,text in enumerate([
        'Excel finalization pending: native pivot status will be written here.',
        'Manual monthly pivot: click MonthlySummary > Insert > PivotTable > Existing Worksheet > Monthly Sales!I5. Rows = Month; Values = Sum of Revenue GBP and Sum of Completed orders; format month as mmm yyyy and GBP with two decimals.',
        'Manual country pivot: click CountrySummary > Insert > PivotTable > Existing Worksheet > Country Performance!I5. Rows = Country; Values = Sum of Revenue GBP and Sum of Completed orders; format GBP, sort revenue descending.',
        'These pivots summarize already verified aggregate rows. They do not provide transaction drill-through or new distinct-customer calculations. Country/month order sums are valid because invoice attribution checks passed.',
        'The other formatted blocks are summary tables, not PivotTables. For a new transaction-level PivotTable use the separate manual guide and the full cleaned source, never just a top-ten table.',
        'Repeat rate denominator: all known IDs with eligible sales in this observation window. Anonymous sales remain in overall sales and are excluded only from customer-specific measures.'
    ],44):
        note(sources,row,text,height=44)
    sources.print_area = 'A1:G50'
    sources.print_title_rows = '1:2'
    # Keep the common B6 split; changing split axes after initialization can
    # leave incompatible pane selections in Excel's serialized view metadata.

    # Monthly data and formulas use the verified aggregate grain.
    monthly = book['Monthly Sales']
    monthly.column_dimensions['A'].width = 15
    monthly.column_dimensions['G'].width = 47
    note(monthly,3,'Completed-sales revenue is GBP, not net revenue. Dec 2011 is partial through 9 Dec; Dec 2010 is a boundary with completeness unverified.',height=32,fill=AMBER)
    rows = []
    status_labels = {'partial_end':'PARTIAL - through 09 Dec 2011',
                     'boundary_start_day_1_unverified':'BOUNDARY - starts day 1; completeness unverified',
                     'interior_calendar_month':'Interior calendar month'}
    for r in frames['monthly'].itertuples(index=False):
        rows.append([datetime.strptime(r.Month,'%Y-%m'), Decimal(r.RevenueGBP),int(r.CompletedOrders),
                     None,None,None,status_labels[r.CoverageStatus]])
    table(monthly,5,['Month','Revenue GBP','Completed orders','AOV GBP','Observed MoM %','Comparable MoM %','Coverage note'],rows,'MonthlySummary')
    for i,r in enumerate(frames['monthly'].itertuples(index=False),6):
        monthly.cell(i,1).number_format = 'mmm yyyy'
        monthly.cell(i,2).number_format = GBP
        monthly.cell(i,3).number_format = INTEGER
        monthly.cell(i,4,f'=ROUND(B{i}/C{i},2)').number_format = GBP
        for c in [5,6]: monthly.cell(i,c).number_format=PERCENT
        monthly.cell(i,7).alignment=Alignment(wrap_text=True,vertical='center')
        monthly.row_dimensions[i].height = 30 if r.CoverageStatus!='interior_calendar_month' else 23
        expect(monthly,f'B{i}',r.RevenueGBP)
        expect(monthly,f'C{i}',r.CompletedOrders)
        expect(monthly,f'D{i}',(Decimal(r.RevenueGBP)/int(r.CompletedOrders)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP))
        if i>6:
            monthly.cell(i,5,f'=IFERROR(B{i}/B{i-1}-1,"")')
            monthly.cell(i,6,f'=IF(AND(G{i}="Interior calendar month",G{i-1}="Interior calendar month"),E{i},"")')
            expect(monthly,f'E{i}',Decimal(r.ObservedMoMPct)/100,'percent')
        if r.ComparableMoMPct:
            expect(monthly,f'F{i}',Decimal(r.ComparableMoMPct)/100,'percent')
        else:
            expect(monthly,f'F{i}','','blank')
        if r.CoverageStatus!='interior_calendar_month':
            for c in monthly[i][:7]: c.fill=PatternFill('solid',fgColor=AMBER)
    monthly['A20']='Full-dataset total'
    monthly['B20']='=SUM(B6:B18)'; monthly['B20'].number_format=GBP
    monthly['C20']='=SUM(C6:C18)'; monthly['C20'].number_format=INTEGER
    expect(monthly,'B20',kpi.RevenueGBP); expect(monthly,'C20',kpi.CompletedOrders)
    monthly.merge_cells('I2:L3'); monthly['I2']='Native PivotTable: monthly summary'
    monthly['I2'].alignment=Alignment(wrap_text=True)
    monthly.print_area='A1:L45'

    products=book['Top Products']
    for col,width in {'A':8,'B':15,'C':47,'D':16,'E':22,'F':3,'G':17}.items(): products.column_dimensions[col].width=width
    note(products,3,'Separate top tens from the verified Python/MariaDB results. Exact StockCode is the key; descriptions are observed labels, not categories.',height=32)
    for name,start in [('top_products_revenue',5),('top_products_quantity',23)]:
        data=[[int(r.Rank),r.StockCode,r.Description,int(r.UnitsSold),Decimal(r.RevenueGBP)] for r in frames[name].itertuples(index=False)]
        table(products,start,['Rank','StockCode','Description','Units sold','Revenue GBP'],data,'ProductsByRevenue' if start==5 else 'ProductsByQuantity')
        for i,r in enumerate(frames[name].itertuples(index=False),start+1):
            products.cell(i,2).number_format='@'
            products.cell(i,3).alignment=Alignment(wrap_text=True,vertical='center')
            products.cell(i,4).number_format=INTEGER; products.cell(i,5).number_format=GBP
            products.row_dimensions[i].height=30
            expect(products,f'B{i}',r.StockCode,'text'); expect(products,f'D{i}',r.UnitsSold); expect(products,f'E{i}',r.RevenueGBP)
    products.print_area='A1:L40'

    country=book['Country Performance']
    for col,width in {'A':25,'B':22,'C':19,'D':20,'E':17,'F':3,'G':16,'H':3}.items(): country.column_dimensions[col].width=width
    note(country,3,'All countries remain in the table. Chart excludes the UK so smaller markets are readable; the title states this explicitly. Anonymous sales remain included.',height=34)
    country_rows=[[r.Country,Decimal(r.RevenueGBP),int(r.CompletedOrders),None,None] for r in frames['countries'].itertuples(index=False)]
    table(country,5,['Country','Revenue GBP','Completed orders','AOV GBP','Revenue share'],country_rows,'CountrySummary')
    for i,r in enumerate(frames['countries'].itertuples(index=False),6):
        country.cell(i,2).number_format=GBP; country.cell(i,3).number_format=INTEGER
        country.cell(i,4,f'=ROUND(B{i}/C{i},2)').number_format=GBP
        country.cell(i,5,f'=B{i}/\'Overview\'!$B$7').number_format=PERCENT
        country.row_dimensions[i].height=19
        expect(country,f'A{i}',r.Country,'text'); expect(country,f'B{i}',r.RevenueGBP); expect(country,f'C{i}',r.CompletedOrders)
        expect(country,f'D{i}',(Decimal(r.RevenueGBP)/int(r.CompletedOrders)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP))
        expect(country,f'E{i}',Decimal(r.RevenueGBP)/Decimal(kpi.RevenueGBP),'percent')
    assert frames['countries'].Country.iloc[0]=='United Kingdom'
    country.merge_cells('I2:L3'); country['I2']='Native PivotTable: country summary'
    country['I2'].alignment=Alignment(wrap_text=True)
    country.print_area='A1:L77'

    customer=book['Customer Analysis']
    note(customer,3,'Repeat customer = known ID with at least two distinct completed invoices in this observation window. Anonymous sales are not one fabricated customer.',height=34,fill=AMBER)
    table(customer,5,['Customer metric','Value','Definition'],[
        ['Known customers',None,'Distinct nonmissing eligible CustomerID'],
        ['Known-customer denominator',int(repeat.KnownCustomers),'Known customers with at least one completed order'],
        ['Repeat customers',int(repeat.RepeatCustomers),'At least two distinct completed invoices'],
        ['One-order customers',int(repeat.OneOrderCustomers),'Exactly one completed invoice'],
        ['Repeat-customer rate',None,'Repeat customers / known-customer denominator'],
        ['Known-customer orders',int(repeat.KnownCustomerOrders),'Distinct eligible invoices for known IDs'],
        ['Known-customer revenue GBP',Decimal(repeat.KnownCustomerRevenueGBP),'Anonymous sales excluded only from this customer metric'],
        ['Eligible anonymous lines',cleaning['retained_missing_customer_rows'],'Remain in overall sales and orders']], 'CustomerMetrics')
    customer['B6']='=B7'; customer['B10']='=B8/B7'
    customer['B6'].number_format=INTEGER
    customer['B10'].number_format=PERCENT; customer['B12'].number_format=GBP
    customer.column_dimensions['A'].width=35; customer.column_dimensions['C'].width=58
    for row in range(6,14):
        customer.cell(row,3).alignment=Alignment(wrap_text=True,vertical='center')
        customer.row_dimensions[row].height=30
    for cell,value in [('B6',repeat.KnownCustomers),('B7',repeat.KnownCustomers),('B8',repeat.RepeatCustomers),('B9',repeat.OneOrderCustomers),('B11',repeat.KnownCustomerOrders),('B12',repeat.KnownCustomerRevenueGBP),('B13',cleaning['retained_missing_customer_rows'])]: expect(customer,cell,value)
    expect(customer,'B10',Decimal(repeat.RepeatCustomerPct)/100,'percent')
    section(customer,16,'Purchase frequency of known customers')
    bands=[('1 order',lambda s:s.eq(1)),('2 orders',lambda s:s.eq(2)),('3-4 orders',lambda s:s.between(3,4)),
           ('5-9 orders',lambda s:s.between(5,9)),('10+ orders',lambda s:s.ge(10))]
    bandrows=[]
    for label,mask in bands:
        group=customers[mask(customers.CompletedOrders)]
        bandrows.append([label,len(group),None,sum(group.RevenueGBP,Decimal(0)),int(group.CompletedOrders.sum())])
    assert sum(r[1] for r in bandrows)==int(repeat.KnownCustomers)
    table(customer,18,['Completed-order band','Known customers','Share of known IDs','Revenue GBP','Completed orders'],bandrows,'CustomerFrequency')
    for row in range(19,24):
        customer.cell(row,3,f'=B{row}/$B$7').number_format=PERCENT
        customer.cell(row,4).number_format=GBP
    customer['A25']='Reconciled known IDs'; customer['B25']='=SUM(B19:B23)'
    customer['B25'].number_format=INTEGER
    expect(customer,'B25',repeat.KnownCustomers)
    note(customer,28,'The rate is 2,829 / 4,334 = 65.27% after display rounding. It is not lifetime retention, a cohort retention rate or a causal effect.',height=34)
    note(customer,30,'Frequency bands are derived solely from each known customer\'s verified distinct completed-order count. No product categories or customer identities were inferred.',height=34)
    customer.print_area='A1:G33'

    quality=book['Data Quality']
    quality.column_dimensions['A'].width=38
    note(quality,3,'Sequential rule exclusions are disjoint; issue flags overlap. Signed excluded line values are not lost revenue or net returns.',height=34,fill=AMBER)
    stage_rows=[[r['rule'],r['before'],r['excluded'],r['after'],Decimal(r['excluded_signed_GBP']),Decimal(r['excluded_absolute_GBP'])] for r in cleaning['stages']]
    table(quality,5,['First applicable exclusion rule','Before rows','Excluded rows','After rows','Signed value GBP','Absolute value GBP'],stage_rows,'CleaningExclusions')
    for row in range(6,12):
        quality.cell(row,1).alignment=Alignment(wrap_text=True,vertical='center')
        quality.row_dimensions[row].height=29
        quality.cell(row,5).number_format=GBP; quality.cell(row,6).number_format=GBP
    quality['A13']='Total excluded lines'; quality['C13']='=SUM(C6:C11)'
    quality['C13'].number_format=INTEGER
    expect(quality,'C13',cleaning['excluded_rows'])
    f=cleaning['flag_counts']
    flagrows=[['Raw lines missing CustomerID',f['Flag_missing_CustomerID'],None,'All 541,909 source lines'],
              ['Eligible anonymous lines',cleaning['retained_missing_customer_rows'],None,'All 522,540 eligible lines'],
              ['Missing descriptions',f['Flag_missing_Description'],None,'All source lines'],
              ['Cancellation lines',f['Flag_cancellation'],None,'All source lines'],
              ['Negative quantities',f['Flag_negative_quantity'],None,'All source lines'],
              ['Nonpositive prices',f['Flag_nonpositive_price'],None,'All source lines'],
              ['Exact duplicate extras (source)',f['Flag_exact_duplicate_extra'],None,'All source lines; overlaps earlier exclusions'],
              ['Conversion failures',sum(f[c] for c in f if c.startswith('Flag_conversion_failure_')),None,'Actual conversions; zero for this input']]
    table(quality,16,['Overlapping issue / retained group','Lines','Share of denominator','Denominator / interpretation'],flagrows,'QualityFlags')
    quality.column_dimensions['D'].width=42
    for row in range(17,25):
        ref=references['Completed-sales rows'] if row==18 else references['Raw rows']
        quality.cell(row,3,f'=B{row}/{ref}').number_format=PERCENT
        quality.cell(row,4).alignment=Alignment(wrap_text=True,vertical='center')
        quality.row_dimensions[row].height=30
        expect(quality,f'B{row}',flagrows[row-17][1])
    section(quality,27,'Limitations and interpretation')
    for row,text in enumerate([
        'Completed-sales revenue is not net revenue or profit. Cancellations, negative adjustments, nonpositive prices and reviewed non-product items are excluded.',
        '5,268 source duplicate extras were flagged. The final duplicate rule removed 5,221 otherwise eligible lines worth GBP 24,213.74; 47 extras had already been excluded by earlier rules.',
        'Exact-duplicate removal is an assumption: identical lines can be legitimate. InvoiceNo is never deduplicated by itself. The source workbook is unchanged.',
        'Missing customer IDs restrict customer metrics but do not invalidate sales. Do not fill all unknown IDs with one customer. Missing descriptions alone do not invalidate recognized merchandise.',
        'Dec 2010: starts day 1 but completeness unverified. Dec 2011: partial through 9 Dec. Do not compare boundary months with full months without qualification.',
        'Case-distinct product codes remain separate. Large transactions remain and may affect rankings. Descriptions vary; no categories, costs, profit, return-to-sale matches or causal explanations are invented.'
    ],28): note(quality,row,text,height=38)
    quality.print_area='A1:G35'

    overview=book['Overview']
    note(overview,3,'COMPLETED SALES, GBP - excludes cancellations and negative adjustments; not net revenue or profit. Snapshot: 01 Dec 2010 to 09 Dec 2011.',height=36,fill=AMBER)
    table(overview,6,['Metric','Workbook value','Python reference','Reconciliation'],[
        ['Completed-sales revenue GBP',None,Decimal(kpi.RevenueGBP),None],
        ['Distinct completed orders',None,int(kpi.CompletedOrders),None],
        ['Units sold',None,int(kpi.UnitsSold),None],
        ['Average order value GBP',None,Decimal(kpi.AverageOrderValueGBP),None],
        ['Known unique customers',None,int(kpi.KnownUniqueCustomers),None],
        ['Repeat customers',None,int(repeat.RepeatCustomers),None],
        ['Repeat-customer rate',None,Decimal(repeat.RepeatCustomerPct)/100,None]],'OverviewKPIs')
    formulas={7:"=ROUND(SUM('Monthly Sales'!B6:B18),2)",8:"=SUM('Monthly Sales'!C6:C18)",
              9:'='+references['Units sold'],10:'=ROUND(B7/B8,2)',11:"='Customer Analysis'!B7",
              12:"='Customer Analysis'!B8",13:"='Customer Analysis'!B10"}
    for row,formula in formulas.items():
        overview.cell(row,2,formula)
        overview.cell(row,4,f'=IF(ROUND(B{row}-C{row},2)=0,"PASS","CHECK")' if row!=13 else f'=IF(ROUND(B{row}*100,2)=ROUND(C{row}*100,2),"PASS","CHECK")')
        overview.row_dimensions[row].height=28
        for col in [2,3]: overview.cell(row,col).number_format=PERCENT if row==13 else (GBP if row in [7,10] else INTEGER)
        expect(overview,f'B{row}',overview.cell(row,3).value,'percent' if row==13 else 'number')
        expect(overview,f'D{row}','PASS','text')
    overview['B9'].comment=Comment('Units are the verified whole-dataset Python/SQL total. The top-ten tables alone cannot reconstruct all units.','Project')
    overview.conditional_formatting.add('D7:D13',CellIsRule(operator='equal',formula=['"CHECK"'],fill=PatternFill('solid',fgColor='FFC7CE'),font=Font(color='9C0006')))
    section(overview,16,'Metric definitions')
    definitions=[('Revenue','Sum of eligible line-rounded integer pence / 100. Per line: Quantity x UnitPrice, ROUND_HALF_UP to GBP 0.01.'),
                 ('Completed order','Distinct InvoiceNo containing eligible merchandise; an invoice can have multiple product lines.'),
                 ('Units / AOV','Units = sum of eligible Quantity. AOV = completed-sales revenue / distinct completed orders.'),
                 ('Known customer','Distinct nonmissing CustomerID with eligible sales. Anonymous sales remain in overall revenue and orders.'),
                 ('Repeat rate','Known IDs with at least two distinct completed invoices / all known IDs with eligible sales. Denominator = 4,334.'),
                 ('Monthly change','Observed MoM is a diagnostic. Comparable MoM excludes boundary pairs; December 2011 is partial.'),
                 ('Product / country','Products group by exact StockCode; no inferred categories. Country is the observed source label.')]
    for row,(label,text) in enumerate(definitions,17): note(overview,row,label+': '+text,height=32)
    section(overview,26,'Business questions and scope')
    for row,text in enumerate([
        'How much eligible merchandise was sold, through how many orders, and at what average value?',
        'Which products and markets lead sales, and how do calendar-covered monthly totals vary?',
        'How many known customers return within the observed window, and how do data-quality decisions affect interpretation?',
        'Completed view: positive quantities/prices, non-cancellation reviewed product lines, with later exact duplicate copies excluded. This is a dataset-specific portfolio analysis, not a claim of business impact.',
        'Native monthly/country PivotTables are based on embedded aggregate summaries. Summary-table filters and pivots do not change full-dataset KPIs. See Sources for regeneration and manual steps.'
    ],27): note(overview,row,text,height=34)
    note(overview,34,'Evidence: all seven MariaDB 10.4.32 result tables matched Python. This workbook reconciles the same outputs; no cleaning or metric definitions were changed.',height=34,fill=LIGHT)
    overview.print_area='A1:G36'; overview.freeze_panes='B7'

    for sheet in book:
        for row in sheet:
            for c in row:
                if c.value is None: continue
                if not c.font.name: c.font=Font(name='Calibri',size=11,color=NAVY)
                if c.data_type=='f':
                    c.font=Font(name='Calibri',size=11,color='087E8B')
                if isinstance(c.value,(int,float,Decimal)) and c.number_format=='General': c.number_format=INTEGER
        sheet.sheet_properties.outlinePr.summaryRight=False
        sheet.print_title_rows='1:2'
        sheet.page_setup.fitToHeight=1 if sheet.title!='Sources' and sheet.title!='Country Performance' else 2
    for row in range(6,44): country.row_dimensions[row].height=16
    country.print_area='A1:L71'
    monthly.print_area='A1:L40'
    products.print_area='A1:L35'
    draft=work/'draft.xlsx'
    book.save(draft)
    plan={'analysis':analysis.relative_to(ROOT).as_posix(),'cleaning':cleaning_file.relative_to(ROOT).as_posix(),
          'sql_evidence':args.sql_evidence.resolve().relative_to(ROOT).as_posix(),
          'expected_cells':expected,'source_hashes':{name:sha256(analysis/f'{name}.csv') for name in KEYS}}
    (work/'validation_plan.json').write_text(json.dumps(plan,indent=2),encoding='utf-8')
    return work


def verify(work):
    target=work/'ready.xlsx'
    values=load_workbook(target,data_only=True)
    formulas=load_workbook(target,data_only=False)
    assert values.sheetnames==SHEETS
    plan=json.loads((work/'validation_plan.json').read_text())
    for check in plan['expected_cells']:
        actual=values[check['sheet']][check['cell']].value
        want=check['expected']
        if check['kind']=='blank': assert actual in [None,''],check
        elif check['kind']=='text': assert str(actual)==want,(check,actual)
        elif check['kind']=='percent':
            assert (Decimal(str(actual))*100).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)==(Decimal(want)*100).quantize(Decimal('.01'),rounding=ROUND_HALF_UP),(check,actual)
        else: assert Decimal(str(actual)).quantize(Decimal('.01'))==Decimal(want).quantize(Decimal('.01')),(check,actual)
    formula_count=0
    for s in formulas:
        assert s.freeze_panes
        for row in s:
            for c in row:
                if c.data_type=='f':
                    formula_count+=1
                    cached=values[s.title][c.coordinate]
                    assert cached.data_type!='e',(s.title,c.coordinate,cached.value)
                    assert '#REF!' not in c.value
    info=json.loads((work/'excel_engine.json').read_text(encoding='utf-8-sig'))
    assert info['PivotTables']==2 and info['Charts']==4,info
    assert sum(len(s._pivots) for s in formulas)==2
    assert sum(len(s._charts) for s in formulas)==4
    for name in ['Monthly Sales','Country Performance']:
        totals=[row for row in values[name].iter_rows(min_col=9,max_col=11)
                if row[0].value=='Grand Total']
        assert len(totals)==1,(name,'Native PivotTable grand total missing')
        assert Decimal(str(totals[0][1].value))==Decimal('10247219.32')
        assert totals[0][2].value==19773
    report={'status':'PASS','checked_reference_cells':len(plan['expected_cells']),
            'formula_count':formula_count,'formula_errors':0,'native_pivottables':2,'native_charts':4,
            'engine':info,'sheets':SHEETS,'source_hashes':plan['source_hashes'],
            'scope':'Excel calculated caches compared to verified Python outputs; SQL evidence validated. Visual previews require separate inspection.'}
    (work/'verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    return report


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--analysis',type=Path,default=ROOT/'reports/analysis_20261017T165141_531793Z')
    p.add_argument('--sql-evidence',type=Path,default=ROOT/'reports/sql_20261017T171530_730533Z_reconciliation.json')
    p.add_argument('--output',type=Path,default=ROOT/'reports/Retail_Sales_Analysis.xlsx')
    p.add_argument('--replace',action='store_true')
    p.add_argument('--prepare-only',action='store_true',help='Prepare draft; finalize in Excel separately, then use --verify-work')
    p.add_argument('--verify-work',type=Path,help='Verify an already Excel-finalized staging folder and install final workbook')
    args=p.parse_args()
    if args.output.exists() and not args.replace: raise FileExistsError('Output exists; --replace archives it before replacement')
    work=args.verify_work.resolve() if args.verify_work else build(args)
    print('Staging folder:',work,flush=True)
    if args.prepare_only: return
    if not args.verify_work:
        subprocess.run(['powershell','-NoProfile','-ExecutionPolicy','Bypass','-File',str(ROOT/'src/finalize_excel.ps1'),'-WorkDir',str(work)],check=True)
    report=verify(work)
    if args.output.exists():
        backup=ROOT/'reports/excel_qa/backups'/datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S_%fZ.xlsx')
        backup.parent.mkdir(parents=True,exist_ok=True); shutil.copy2(args.output,backup)
    args.output.parent.mkdir(parents=True,exist_ok=True)
    os.replace(work/'ready.xlsx',args.output)
    (ROOT/'reports/excel_validation.json').write_text(json.dumps({**report,'workbook':args.output.resolve().relative_to(ROOT).as_posix(),
        'workbook_sha256':sha256(args.output),'qa_folder':work.relative_to(ROOT).as_posix()},indent=2),encoding='utf-8')
    print('Created:',args.output)
    print(json.dumps(report,indent=2))


if __name__=='__main__': main()
